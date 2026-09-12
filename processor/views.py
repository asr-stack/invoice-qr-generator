import os
import re
import shutil
import tempfile
import uuid
import zipfile
from datetime import timedelta
from io import BytesIO
import zipfile as pyzipfile

from django.conf import settings
from django.http import FileResponse, JsonResponse
from django.shortcuts import render
from django.utils import timezone
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_POST
from openpyxl import load_workbook
from docx import Document

from .utils import (
    extract_invoice_number,
    generate_qr,
    insert_qr_into_docx,
    read_excel,
)

ALLOWED_EXCEL_EXTENSIONS = {".xlsx", ".xls"}
ALLOWED_TEMPLATE_EXTENSIONS = {".docx"}


def _session_batch_dir(request):
    batch_id = request.session.get('batch_id')
    if not batch_id:
        batch_id = uuid.uuid4().hex
        request.session['batch_id'] = batch_id

    batch_root = os.path.abspath(os.path.join(settings.MEDIA_ROOT, 'temp', batch_id))
    processed_dir = os.path.join(batch_root, 'processed')
    os.makedirs(processed_dir, exist_ok=True)
    return batch_id, batch_root, processed_dir


def _cleanup_stale_batches():
    temp_root = os.path.abspath(os.path.join(settings.MEDIA_ROOT, 'temp'))
    if not os.path.isdir(temp_root):
        return

    cutoff = timezone.now() - timedelta(hours=getattr(settings, 'TEMP_FILE_RETENTION_HOURS', 1))
    cutoff_ts = cutoff.timestamp()

    for entry in os.scandir(temp_root):
        if not entry.is_dir():
            continue
        try:
            if os.path.getmtime(entry.path) < cutoff_ts:
                shutil.rmtree(entry.path, ignore_errors=True)
        except OSError:
            continue

    # Also remove any empty or orphaned batch directories left behind during failed flows.
    for entry in os.scandir(temp_root):
        if not entry.is_dir():
            continue
        try:
            if not os.listdir(entry.path):
                shutil.rmtree(entry.path, ignore_errors=True)
        except OSError:
            continue


def _validate_file_and_size(uploaded_file, allowed_extensions, label):
    if uploaded_file is None:
        raise ValueError(f"Please upload a valid {label} file.")

    filename = (uploaded_file.name or '').strip()
    if not filename:
        raise ValueError(f"Please upload a valid {label} file.")

    if uploaded_file.size > getattr(settings, 'MAX_UPLOAD_SIZE', 10 * 1024 * 1024):
        raise ValueError(f"{label} file exceeds the allowed size limit.")

    _, ext = os.path.splitext(filename.lower())
    if ext not in allowed_extensions:
        raise ValueError(f"Unsupported {label} format. Please upload only {', '.join(sorted(allowed_extensions))}.")

    return filename


def _is_valid_xlsx_file(file_obj):
    try:
        file_obj.seek(0)
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
        workbook.close()
        file_obj.seek(0)
        return True
    except Exception:
        return False


def _is_valid_docx_file(file_obj):
    try:
        file_obj.seek(0)
        file_obj.read(4)
        file_obj.seek(0)
        with pyzipfile.ZipFile(file_obj, 'r') as zf:
            names = set(zf.namelist())
            required = {'[Content_Types].xml', 'word/document.xml'}
            if not required.issubset(names):
                return False

        file_obj.seek(0)
        Document(file_obj)
        file_obj.seek(0)
        return True
    except Exception:
        return False


def _safe_file_path(base_dir, relative_name):
    if not relative_name:
        raise ValueError('Invalid file request')

    if relative_name.startswith(('/', '\\')):
        raise ValueError('Invalid file request')

    if '..' in relative_name.split('/') or '..' in relative_name.split('\\'):
        raise ValueError('Invalid file request')

    candidate = os.path.abspath(os.path.join(base_dir, os.path.basename(relative_name)))
    if os.path.commonpath([base_dir, candidate]) != base_dir:
        raise ValueError('Invalid file request')

    return candidate


def _session_downloads(request):
    return request.session.setdefault('batch_downloads', {})


def _make_download_id(file_name, file_path):
    return uuid.uuid4().hex


def home(request):
    _cleanup_stale_batches()
    return render(request, 'index.html')


@require_POST
@csrf_protect
def upload_files(request):
    _cleanup_stale_batches()

    batch_id, batch_root, processed_dir = _session_batch_dir(request)
    data_file = request.FILES.get('data_file')
    templates = request.FILES.getlist('templates')

    print("Excel received:", data_file.name if data_file else None)
    print("Word files received:", [f.name for f in templates])

    if not data_file:
        return JsonResponse({'success': False, 'error': 'Please upload an Excel file.'}, status=400)

    if not templates:
        return JsonResponse({'success': False, 'error': 'Please upload at least one Word document.'}, status=400)

    try:
        _validate_file_and_size(data_file, ALLOWED_EXCEL_EXTENSIONS, 'Excel')
    except ValueError as exc:
        return JsonResponse({'success': False, 'error': str(exc)}, status=400)

    for template in templates:
        try:
            _validate_file_and_size(template, ALLOWED_TEMPLATE_EXTENSIONS, 'Word')
        except ValueError as exc:
            return JsonResponse({'success': False, 'error': str(exc)}, status=400)

    try:
        if not _is_valid_xlsx_file(data_file):
            return JsonResponse({'success': False, 'error': '❌ Invalid Excel file. Please upload a valid .xlsx or .xls file.'}, status=400)
        excel_map = read_excel(data_file)
    except Exception:
        return JsonResponse({'success': False, 'error': '❌ Invalid Excel file. Please upload a valid .xlsx or .xls file.'}, status=400)

    results = []

    for template in templates:
        temp_docx_path = None
        try:
            print(f"📄 Processing: {template.name}")

            with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp:
                for chunk in template.chunks():
                    tmp.write(chunk)
                temp_docx_path = tmp.name

            with open(temp_docx_path, 'rb') as f:
                if not _is_valid_docx_file(f):
                    results.append({'file': template.name, 'status': '❌ Invalid Word document. Please upload a valid .docx file.'})
                    continue

            match = re.search(r'Invoice_(\d+)', template.name)
            if match:
                invoice_no = match.group(1)
            else:
                invoice_no = extract_invoice_number(temp_docx_path)

            if invoice_no is None:
                results.append({'file': template.name, 'status': 'Invoice not found'})
                continue

            print(f"Invoice from filename: {invoice_no}")

            data = excel_map.get(invoice_no)
            if data is None:
                results.append({'file': template.name, 'status': f'No data for invoice {invoice_no}'})
                continue

            print(f"Matched invoice data: {data}")

            irn = data['irn']
            ack_no = data['ack_no']
            ack_date = data['ack_date']
            doc_no = data.get('doc_no') or invoice_no
            doc_date = data.get('doc_date')
            inv_value = data.get('inv_value')

            qr_image = generate_qr(irn, ack_no, ack_date, doc_no, doc_date, inv_value)
            final_docx = insert_qr_into_docx(
                temp_docx_path,
                qr_image,
                irn,
                ack_no,
                ack_date,
                template.name,
                output_dir=processed_dir,
            )

            download_id = _make_download_id(template.name, final_docx)
            download_map = _session_downloads(request)
            download_map[download_id] = os.path.basename(final_docx)
            request.session['batch_id'] = batch_id
            request.session.modified = True

            results.append({
                'file': os.path.basename(final_docx),
                'status': 'Success',
                'download_id': download_id,
            })
        except Exception:
            results.append({'file': template.name, 'status': 'Unable to process this document.'})
        finally:
            if temp_docx_path and os.path.exists(temp_docx_path):
                try:
                    os.unlink(temp_docx_path)
                except OSError:
                    pass

    return JsonResponse({'results': results})


def download_file(request, download_id):
    if not request.session.get('batch_id'):
        return JsonResponse({'error': 'Invalid file request'}, status=400)

    batch_id = request.session['batch_id']
    batch_dir = os.path.abspath(os.path.join(settings.MEDIA_ROOT, 'temp', batch_id, 'processed'))

    if not os.path.isdir(batch_dir):
        return JsonResponse({'error': 'File not found'}, status=404)

    file_name = _session_downloads(request).get(download_id)
    if not file_name:
        return JsonResponse({'error': 'Invalid file request'}, status=400)

    try:
        target_path = _safe_file_path(batch_dir, file_name)
    except ValueError:
        return JsonResponse({'error': 'Invalid file request'}, status=400)

    if not os.path.isfile(target_path) or not target_path.lower().endswith('.docx'):
        return JsonResponse({'error': 'File not found'}, status=404)

    return FileResponse(open(target_path, 'rb'), as_attachment=True, filename=os.path.basename(target_path))


def download_all(request):
    batch_id = request.session.get('batch_id')
    if not batch_id:
        return JsonResponse({'error': 'No processed files available'}, status=404)

    batch_dir = os.path.abspath(os.path.join(settings.MEDIA_ROOT, 'temp', batch_id, 'processed'))
    if not os.path.isdir(batch_dir):
        return JsonResponse({'error': 'No processed files available'}, status=404)

    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        for filename in sorted(os.listdir(batch_dir)):
            if not filename.lower().endswith('.docx'):
                continue
            file_path = os.path.join(batch_dir, filename)
            if os.path.isfile(file_path):
                zip_file.write(file_path, arcname=filename)

    zip_buffer.seek(0)
    return FileResponse(zip_buffer, as_attachment=True, filename='Processed_Invoices.zip')


@require_POST
@csrf_protect
def clear_all(request):
    batch_id = request.session.get('batch_id')
    if batch_id:
        batch_root = os.path.abspath(os.path.join(settings.MEDIA_ROOT, 'temp', batch_id))
        if os.path.isdir(batch_root):
            shutil.rmtree(batch_root, ignore_errors=True)

    request.session.pop('batch_id', None)
    request.session.pop('batch_downloads', None)
    request.session.flush()
    request.session.modified = True
    return JsonResponse({'message': 'Current batch cleared and all temporary files deleted.'})
